"""Generate the official SIH idea-presentation deck (4 slides) for SIH26162.

Format mirrors the standard SIH template:
  Slide 1 — Basic Details of the Team and Problem Statement
  Slide 2 — Problem + Idea/Solution + Technology Stack + Flow Chart
  Slide 3 — Idea/Approach Details (Use Cases, Dependencies, Channels, Revenue)
  Slide 4 — Team Member Details

    ./.venv/bin/python make_sih_ppt.py

Edit TEAM / MENTORS / FILL placeholders below, re-run, done.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---------------------------------------------------------------- palette --
DARK = RGBColor(0x1A, 0x1A, 0x2E)
RED = RGBColor(0xE6, 0x39, 0x46)
ORANGE = RGBColor(0xF7, 0x7F, 0x00)
TEXT = RGBColor(0x2B, 0x2B, 0x2B)
MUTED = RGBColor(0x6C, 0x75, 0x7D)
NAVY = RGBColor(0x1D, 0x35, 0x57)
SLATE = RGBColor(0x45, 0x7B, 0x9D)
PURPLE = RGBColor(0x7B, 0x2C, 0xBF)
TEAL = RGBColor(0x2A, 0x9D, 0x8F)
BOX_RED = RGBColor(0xFF, 0xF4, 0xEF)
BOX_GRAY = RGBColor(0xF5, 0xF6, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "Calibri"

# ------------------------------------------------------- editable content --
TEAM = {
    "ministry": "National Technical Research Organisation (NTRO)",
    "ps_code": "SIH26162",
    "ps_title": ("AI-Based Detection and Classification of Industrial Fires and "
                 "Persistent Thermal Sources Using NASA FIRMS, OSM & Satellite Data"),
    "theme": "Disaster Management (Software Edition)",
    "team_name": "AgniNetra",
    "idea_title": "AgniNetra - AI-Powered Detection & Classification of Industrial Fires "
                  "and Persistent Thermal Sources",
    "leader": "FILL: Leader Name",
    "institute_code": "FILL: AISHE code",
    "institute": "FILL: Institute Name",
}
MEMBERS = [  # leader + 5 members
    ("Team Leader Name: ", "FILL: Leader Name"),
    ("Team Member 1 Name: ", "FILL: Name"),
    ("Team Member 2 Name: ", "FILL: Name"),
    ("Team Member 3 Name: ", "FILL: Name"),
    ("Team Member 4 Name: ", "FILL: Name"),
    ("Team Member 5 Name: ", "FILL: Name"),
]
MENTORS = [("Team Mentor 1 Name: ", "FILL: Mentor Name")]

PROBLEM = [
    ("Anonymous hotspots: ", "satellites report every fire as a bare coordinate - a wildfire, "
     "a refinery blaze and a gas flare look identical."),
    ("Slow manual classification: ", "hours-to-days delay lets industrial fires escalate "
     "into catastrophes."),
    ("Invisible chronic polluters: ", "no persistence tracking - gas flares burning for "
     "weeks go unnoticed."),
    ("Data is not GIS-ready: ", "raw satellite feeds give agencies nothing they can drop "
     "into QGIS / Google Earth workflows."),
    ("No automatic alerting: ", "nobody is woken when an industrial fire ignites near a facility."),
]
SOLUTION = [
    ("Real-Time Ingestion: ", "NASA FIRMS (VIIRS S-NPP, NOAA-20, MODIS) hotspots pulled "
     "every 10 minutes across India, de-duplicated into SQLite - 100% real satellite data, "
     "zero synthetic rows."),
    ("AI Classification (XGBoost): ", "each hotspot auto-classified into 6 source types "
     "(wildfire | crop burning | industrial fire | gas flare | mining | other) using 9 "
     "engineered features - brightness, FRP, day/night, persistence, distance to nearest "
     "OSM facility, land-cover. 95.4% test accuracy."),
    ("Geospatial Intelligence: ", "every detection auto-named with its Indian state + nearest "
     "named facility via OSM data & official state boundaries - \"Odisha - 0.8 km from power plant\"."),
    ("Real Alerting Engine: ", "industrial/persistent source + >=85% confidence + <=5 km from "
     "a named facility -> instant Telegram/email alert with Google Maps link (6-h dedup). "
     "96 alerts already fired on live data."),
    ("GIS Dashboard + Export: ", "React dashboard on Google satellite imagery, 10-min live "
     "refresh, filters & trends; one-click GeoJSON / KML / CSV export for QGIS & Google Earth."),
    ("Model Transparency: ", "published validation report - confusion matrix, feature "
     "importance - served via API and rendered inside the dashboard."),
]
TECH_STACK = ("Python 3.13 | FastAPI | SQLite | XGBoost + scikit-learn | httpx | APScheduler | "
              "React (CRA) + Leaflet (Google Satellite layers) | NASA FIRMS API | OpenStreetMap "
              "Overpass | ESA WorldCover | Telegram Bot API | GeoJSON / KML / CSV")
USE_CASES = [
    ("Industrial Safety & Disaster Response: ", "district authorities get a named alert "
     "(\"Chhattisgarh - 0.6 km from refinery works\") within minutes of ignition - not days."),
    ("Environmental Regulation: ", "pollution boards use persistence analytics (detection "
     "count + active days) to catch chronic flare operators, not one-off events."),
    ("Defense / Geospatial Intelligence: ", "GIS-ready KML/GeoJSON feeds plug straight into "
     "NTRO & agency QGIS workflows, including sensitive border regions."),
    ("Agricultural Smoke Management: ", "the crop-burning class + seasonal trends let state "
     "agencies target stubble-burning crackdowns precisely."),
    ("Research & Insurance Forensics: ", "18-column CSV export reproduces the full feature "
     "vector per hotspot for auditable analysis."),
]
DEPENDENCIES = [
    ("NASA FIRMS cadence & availability: ", "mitigated by de-duplication, local DB "
     "persistence, and a live \"last sync\" banner on the dashboard."),
    ("Land-cover coverage gap: ", "current lookup covers ~43% of live points; direct ESA "
     "WorldCover sampling is the scoped fix (next sprint)."),
    ("Label provenance: ", "labels bootstrapped from a rule-based pipeline; a human-labeled "
     "subset + spatial cross-validation are planned hardening steps."),
    ("Class imbalance: ", "gas-flare F1 (66%) limited by a 45-row test class; more flare "
     "training data is the known fix."),
    ("Live-demo connectivity: ", "offline fallback - pre-exported KML in Google Earth + "
     "cached SQLite snapshot."),
]
CHANNELS = ("NTRO & defense geospatial cells | State Pollution Control Boards | SDMA / NDMA "
            "disaster authorities | Forest departments | Refinery & plant HSE teams")
REVENUE = ("SaaS subscription for industrial HSE compliance monitoring | API licensing to "
           "insurers / logistics / ESG analytics | Custom on-prem deployments for government agencies")

FLOW_TOP = [  # (line1, line2, fill)
    ("NASA FIRMS API", "VIIRS - NOAA-20 - MODIS", NAVY),
    ("Ingestion + Dedup", "every 10 minutes", SLATE),
    ("SQLite fires.db", "persistence counters", SLATE),
    ("Enrichment", "facility dist | landcover | state", PURPLE),
    ("XGBoost AI", "6 source classes", RED),
    ("FastAPI REST", "10+ JSON endpoints", ORANGE),
]
FLOW_OUT = [
    ("GIS Dashboard", "React + Google Satellite", TEAL),
    ("Telegram / Email Alerts", "named facility + confidence", TEAL),
    ("GeoJSON | KML | CSV Export", "QGIS + Google Earth", TEAL),
]


# ---------------------------------------------------------------- helpers --
def box(slide, x, y, w, h, fill, line=None, rounded=True, radius=0.05):
    shape_t = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    sp = slide.shapes.add_shape(shape_t, Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    sp.fill.solid()
    sp.fill.fore_color.rgb = fill
    if line is not None:
        sp.line.color.rgb = line
        sp.line.width = Pt(1.2)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def para(tf, parts, size, color=TEXT, bold_all=False, align=None, space_after=None,
         first=False, line_spacing=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    if align is not None:
        p.alignment = align
    if space_after is not None:
        p.space_after = Pt(space_after)
    if line_spacing is not None:
        p.line_spacing = line_spacing
    if isinstance(parts, str):
        parts = [(parts, bold_all)]
    for txt, bold in parts:
        r = p.add_run()
        r.text = txt
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    return p


def title_bar(slide, text):
    bar = box(slide, 0, 0, 20, 0.62, DARK, rounded=False)
    tf = bar.text_frame
    tf.margin_left = Inches(0.35)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = True
    para(tf, text, 19, WHITE, bold_all=True, first=True)
    box(slide, 0, 0.62, 20, 0.05, ORANGE, rounded=False)


def header_tf(slide, x, y, w, h, fill, line, header, header_color):
    sp = box(slide, x, y, w, h, fill, line=line)
    tf = sp.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.18)
    tf.margin_right = Inches(0.18)
    tf.margin_top = Inches(0.12)
    para(tf, header, 15, header_color, bold_all=True, first=True, space_after=6)
    return tf


def flow_box(slide, x, y, w, h, line1, line2, fill):
    sp = box(slide, x, y, w, h, fill, radius=0.10)
    tf = sp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    para(tf, line1, 11, WHITE, bold_all=True, align=PP_ALIGN.CENTER, first=True)
    para(tf, line2, 8.5, WHITE, align=PP_ALIGN.CENTER)


def arrow(slide, x, y, w, h, direction="right"):
    shape_t = MSO_SHAPE.RIGHT_ARROW if direction == "right" else MSO_SHAPE.DOWN_ARROW
    sp = slide.shapes.add_shape(shape_t, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid()
    sp.fill.fore_color.rgb = MUTED
    sp.line.fill.background()
    sp.shadow.inherit = False


# ------------------------------------------------------------------ build --
prs = Presentation()
prs.slide_width = Inches(20)
prs.slide_height = Inches(11.25)
blank = prs.slide_layouts[6]

# ---- Slide 1: Basic details ------------------------------------------------
s = prs.slides.add_slide(blank)
title_bar(s, "Basic Details of the Team and Problem Statement")
sp = box(s, 0.9, 1.15, 18.2, 9.3, BOX_GRAY, line=None)
tf = sp.text_frame
tf.word_wrap = True
tf.margin_left = Inches(0.5)
tf.margin_top = Inches(0.4)
rows = [
    ("Idea Title: ", TEAM["idea_title"], 15),
    ("Ministry/Organization Name: ", TEAM["ministry"], 15),
    ("PS Code: ", TEAM["ps_code"], 15),
    ("Problem Statement Title: ", TEAM["ps_title"], 15),
    ("Theme Name: ", TEAM["theme"], 15),
    ("Team Name: ", TEAM["team_name"], 15),
    ("Team Leader Name: ", TEAM["leader"], 15),
    ("Institute Code (AISHE): ", TEAM["institute_code"], 15),
    ("Institute Name: ", TEAM["institute"], 15),
]
for i, (label, value, size) in enumerate(rows):
    para(tf, [(label, True), (value, False)], size,
         space_after=14, first=(i == 0), line_spacing=1.05)

# ---- Slide 2: Problem + Solution + Tech stack + Flow chart ------------------
s = prs.slides.add_slide(blank)
title_bar(s, "Problem | Idea/Solution | Technology Stack | Flow Chart")

tf = header_tf(s, 0.4, 0.88, 9.3, 4.55, BOX_RED, RED, "Problem :", RED)
for i, (lead, rest) in enumerate(PROBLEM):
    para(tf, [(lead, True), (rest, False)], 11.5, space_after=7, first=(i == 0),
         line_spacing=1.02)

tf = header_tf(s, 9.95, 0.88, 9.65, 4.55, BOX_GRAY, TEAL, "Idea/Solution :", NAVY)
for i, (lead, rest) in enumerate(SOLUTION):
    para(tf, [(lead, True), (rest, False)], 10.5, space_after=5, first=(i == 0),
         line_spacing=1.0)

sp = box(s, 0.4, 5.62, 19.2, 1.0, RGBColor(0xFF, 0xF9, 0xE6), line=ORANGE)
tf = sp.text_frame
tf.word_wrap = True
tf.margin_left = Inches(0.18)
tf.margin_top = Inches(0.08)
para(tf, [("Technology Stack : ", True)], 12.5, color=ORANGE, first=True, space_after=2)
para(tf, TECH_STACK, 11)

para(s.shapes.add_textbox(Inches(0.4), Inches(6.86), Inches(4), Inches(0.4)).text_frame,
     "Flow Chart :", 15, DARK, bold_all=True, first=True)

x0, bw, gap, by, bh = 0.55, 2.72, 0.46, 7.32, 1.15
for i, (l1, l2, fill) in enumerate(FLOW_TOP):
    bx = x0 + i * (bw + gap)
    flow_box(s, bx, by, bw, bh, l1, l2, fill)
    if i < len(FLOW_TOP) - 1:
        arrow(s, bx + bw + 0.04, by + bh / 2 - 0.14, 0.38, 0.28)

# connector from FastAPI down to outputs row
last_cx = x0 + 5 * (bw + gap) + bw / 2
hub_y = by + bh
out_y = 9.0
ow, ogap = 5.9, 0.35
ox0 = (20 - (3 * ow + 2 * ogap)) / 2
box(s, last_cx - 0.02, hub_y, 0.04, 0.18, MUTED, rounded=False)                      # stem
first_cx = ox0 + ow / 2
box(s, first_cx, hub_y + 0.17, last_cx - first_cx, 0.04, MUTED, rounded=False)       # rail
for i, (l1, l2, fill) in enumerate(FLOW_OUT):
    cx = ox0 + i * (ow + ogap) + ow / 2
    arrow(s, cx - 0.15, hub_y + 0.2, 0.3, 0.28, direction="down")
    flow_box(s, ox0 + i * (ow + ogap), out_y, ow, 1.2, l1, l2, fill)

# ---- Slide 3: Approach details ----------------------------------------------
s = prs.slides.add_slide(blank)
title_bar(s, "Idea/Approach Details")

tf = header_tf(s, 0.4, 0.88, 9.5, 5.9, BOX_GRAY, TEAL, "Use Cases :", NAVY)
for i, (lead, rest) in enumerate(USE_CASES):
    para(tf, [(lead, True), (rest, False)], 11.5, space_after=8, first=(i == 0),
         line_spacing=1.03)

tf = header_tf(s, 10.1, 0.88, 9.5, 5.9, BOX_RED, RED, "Dependencies / Show Stoppers :", RED)
for i, (lead, rest) in enumerate(DEPENDENCIES):
    para(tf, [(lead, True), (rest, False)], 11.5, space_after=8, first=(i == 0),
         line_spacing=1.03)

tf = header_tf(s, 0.4, 7.0, 9.5, 1.75, RGBColor(0xE8, 0xF6, 0xF3), TEAL, "Channels :", TEAL)
para(tf, CHANNELS, 11.5, space_after=0, first=True, line_spacing=1.05)

tf = header_tf(s, 10.1, 7.0, 9.5, 1.75, RGBColor(0xE8, 0xF6, 0xF3), TEAL, "Revenue Streams :", TEAL)
para(tf, REVENUE, 11.5, space_after=0, first=True, line_spacing=1.05)

# ---- Slide 4: Team -----------------------------------------------------------
s = prs.slides.add_slide(blank)
title_bar(s, "Team Member Details")
sp = box(s, 0.9, 1.15, 18.2, 9.3, BOX_GRAY)
tf = sp.text_frame
tf.word_wrap = True
tf.margin_left = Inches(0.5)
tf.margin_top = Inches(0.4)
first = True
for label, name in MEMBERS:
    para(tf, [(label, True), (name, False),
              ("        Branch : ", True), ("B.Tech", False),
              ("        Stream : ", True), ("FILL", False),
              ("        Year : ", True), ("FILL", False)],
         15, space_after=14, first=first)
    first = False
for label, name in MENTORS:
    para(tf, [(label, True), (name, False),
              ("        Category : ", True), ("FILL", False),
              ("        Expertise : ", True), ("FILL (e.g. AI/ML, GIS)", False),
              ("        Domain Experience (in years): ", True), ("FILL", False)],
         15, space_after=14)

prs.core_properties.title = "AgniNetra - SIH26162 Idea Presentation"
prs.core_properties.author = "Team AgniNetra"
OUT = "SIH26162_AgniNetra_Idea_Presentation.pptx"
prs.save(OUT)
print(f"saved {OUT} with {len(prs.slides.__iter__.__self__._sldIdLst)} slides")
